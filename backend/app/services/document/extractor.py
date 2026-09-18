import io
from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document as DocxDocument

class DocumentExtractor:
    @staticmethod
    def extract_from_pdf(file_path: Path) -> List[Dict[str, Any]]:
        """
        Extract text page by page from a PDF file.
        Returns a list of dicts: [{"page_number": 1, "text": "..."}]
        """
        pages = []
        try:
            reader = PdfReader(str(file_path))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                # Clean basic whitespace artifacts
                cleaned = "\n".join([line.strip() for line in text.splitlines() if line.strip()])
                pages.append({
                    "page_number": i + 1,
                    "text": cleaned if cleaned else "[Empty page / image content]"
                })
        except Exception as e:
            raise ValueError(f"Failed to extract PDF contents: {str(e)}") from e
        return pages

    @staticmethod
    def extract_from_docx(file_path: Path) -> List[Dict[str, Any]]:
        """
        Extract text paragraph by paragraph / section from a DOCX file.
        Simulates pages by splitting paragraphs into reasonable ~3000-character logical chunks.
        """
        try:
            doc = DocxDocument(str(file_path))
            full_text = []
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text.strip())
            
            # Also extract tables
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_lines.append(" | ".join(row_data))
                if table_lines:
                    full_text.append("\n[Table Data]:\n" + "\n".join(table_lines))

            # Break into logical pages of ~2500-3000 chars or paragraph clusters
            pages = []
            current_page_text = []
            current_len = 0
            page_num = 1

            for para in full_text:
                current_page_text.append(para)
                current_len += len(para)
                if current_len >= 2500:
                    pages.append({
                        "page_number": page_num,
                        "text": "\n\n".join(current_page_text)
                    })
                    page_num += 1
                    current_page_text = []
                    current_len = 0

            if current_page_text:
                pages.append({
                    "page_number": page_num,
                    "text": "\n\n".join(current_page_text)
                })

            if not pages:
                pages.append({"page_number": 1, "text": "[Empty document]"})

            return pages
        except Exception as e:
            raise ValueError(f"Failed to extract DOCX contents: {str(e)}") from e

    @classmethod
    def extract(cls, file_path: Path, file_type: str) -> List[Dict[str, Any]]:
        ext = file_type.lower()
        if not ext.startswith("."):
            ext = f".{ext}"

        if ext == ".pdf":
            return cls.extract_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls.extract_from_docx(file_path)
        elif ext in [".txt", ".md"]:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
            return [{"page_number": 1, "text": text}]
        else:
            raise ValueError(f"Unsupported document type: {ext}. Allowed types: .pdf, .docx, .txt, .md")
