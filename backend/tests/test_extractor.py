import pytest
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from app.services.document.extractor import DocumentExtractor

def test_extract_pdf(tmp_path):
    pdf_path = tmp_path / "test.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.drawString(100, 750, "DocAudit Confidential Legal Agreement Page 1")
    c.drawString(100, 700, "Section 1: Obligations of the Parties")
    c.showPage()
    c.drawString(100, 750, "Page 2: Governing Law and Delaware Jurisdiction")
    c.save()

    pages = DocumentExtractor.extract(pdf_path, ".pdf")
    assert len(pages) == 2
    assert pages[0]["page_number"] == 1
    assert "DocAudit Confidential" in pages[0]["text"]
    assert pages[1]["page_number"] == 2
    assert "Delaware" in pages[1]["text"]

def test_extract_docx(tmp_path):
    docx_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_heading("Financial Statement FY2025", 0)
    doc.add_paragraph("Total recorded revenue is $50,000,000.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "EBITDA"
    table.rows[1].cells[1].text = "12M"
    doc.save(str(docx_path))

    pages = DocumentExtractor.extract(docx_path, ".docx")
    assert len(pages) >= 1
    assert "$50,000,000" in pages[0]["text"]
    assert "EBITDA" in pages[0]["text"]

def test_unsupported_format(tmp_path):
    fake_path = tmp_path / "test.xyz"
    fake_path.write_text("hello")
    with pytest.raises(ValueError, match="Unsupported document type"):
        DocumentExtractor.extract(fake_path, ".xyz")
